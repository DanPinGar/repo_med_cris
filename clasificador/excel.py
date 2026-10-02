"""Lectura y escritura de la tabla de medicamentos."""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ANCHOS = {"Nombre": 30, "API/Comercial": 34, "Principio activo (inglés)": 32, "Descripción": 44,
          "Área": 20, "ATC": 26, "Revisar": 55}
AMARILLO = PatternFill("solid", fgColor="FFF2CC")


def _limpiar(valor) -> str:
    if valor is None:
        return ""
    return " ".join(str(valor).split())


def leer(ruta: Path) -> list[dict]:
    wb = load_workbook(ruta, data_only=True)
    ws = wb["tabla"] if "tabla" in wb.sheetnames else wb.active
    cabeceras = [_limpiar(c.value) for c in ws[1]]
    filas = []
    for valores in ws.iter_rows(min_row=2, values_only=True):
        fila = {cab: _limpiar(v) for cab, v in zip(cabeceras, valores) if cab}
        if fila.get("Nombre"):
            filas.append(fila)
    return filas


def escribir(ruta: Path, filas: list[dict], columnas: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "tabla"
    ws.append(columnas)
    for fila in filas:
        ws.append([fila.get(c, "") for c in columnas])

    ultima_col = get_column_letter(len(columnas))
    ultima_fila = max(ws.max_row, 2)
    tabla = Table(displayName="Medicamentos", ref=f"A1:{ultima_col}{ultima_fila}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)

    for i, columna in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ANCHOS.get(columna, 26)
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"

    if "Revisar" in columnas:
        col = get_column_letter(columnas.index("Revisar") + 1)
        ws.conditional_formatting.add(
            f"A2:{ultima_col}{ultima_fila}", FormulaRule(formula=[f'LEN(${col}2)>0'], fill=AMARILLO)
        )
    wb.save(ruta)


def escribir_informe(ruta: Path, cambios: list[tuple]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "cambios"
    columnas = ["Nombre", "Cambio", "Columna", "Antes", "Después"]
    ws.append(columnas)
    for cambio in cambios:
        ws.append(list(cambio))
    tabla = Table(displayName="Cambios", ref=f"A1:E{max(ws.max_row, 2)}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    for letra, ancho in zip("ABCDE", (30, 34, 26, 45, 45)):
        ws.column_dimensions[letra].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(ruta)
