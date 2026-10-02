"""Lectura de reglas.xlsx y clasificación por ensayo a partir de códigos ATC."""
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .texto import clave, norm


def _lista(valor) -> list[str]:
    return [p for p in re.split(r"[,;\s]+", str(valor or "").upper()) if p]


def _prefijo_mas_largo(codigo: str, tabla: list[tuple[str, str]]) -> str:
    mejor = ("", "")
    for prefijo, texto in tabla:
        if codigo.startswith(prefijo) and len(prefijo) > len(mejor[0]):
            mejor = (prefijo, texto)
    return mejor[1]


@dataclass
class Categoria:
    nombre: str
    incluidos: list[str]
    excluidos: list[str]
    vias_excluidas: list[str]

    def coincide(self, atc: str, vias: list[str]) -> bool:
        if not atc or not any(atc.startswith(p) for p in self.incluidos):
            return False
        if any(atc.startswith(p) for p in self.excluidos):
            return False
        vias_norm = [norm(v) for v in vias]
        return not any(excl in v for excl in self.vias_excluidas for v in vias_norm)


@dataclass
class Ensayo:
    columna: str
    nombre: str
    otros: str
    ninguna: str
    categorias: list[Categoria] = field(default_factory=list)


class Reglas:
    def __init__(self, ruta: Path):
        wb = load_workbook(ruta, read_only=True, data_only=True)

        def filas(hoja):
            datos = list(wb[hoja].iter_rows(min_row=2, values_only=True))
            return [f for f in datos if any(v not in (None, "") for v in f)]

        self.ensayos: list[Ensayo] = []
        por_nombre = {}
        for columna, nombre, otros, ninguna in filas("Ensayos"):
            ensayo = Ensayo(str(columna).strip(), str(nombre).strip(), str(otros or "").strip(), str(ninguna or "-").strip())
            self.ensayos.append(ensayo)
            por_nombre[norm(ensayo.nombre)] = ensayo
        for nombre, categoria, incl, excl, vias, *_ in filas("Categorías"):
            ensayo = por_nombre.get(norm(nombre))
            if ensayo is None:
                raise ValueError(f"reglas.xlsx: la categoría '{categoria}' es de un ensayo desconocido: {nombre}")
            vias_excl = [norm(v) for v in str(vias or "").split(",") if norm(v)]
            ensayo.categorias.append(Categoria(str(categoria).strip(), _lista(incl), _lista(excl), vias_excl))

        config = {norm(p): v for p, v, *_ in filas("Configuración")}
        self.cv_incluidos = _lista(config.get(norm("ATC cardiovasculares")))
        self.cv_excluidos = _lista(config.get(norm("ATC no cardiovasculares")))
        self.sin_descomponer = _lista(config.get(norm("Combinaciones que no se descomponen")))

        self.descripciones = [(str(p).strip().upper(), str(t).strip()) for p, t, *_ in filas("Descripciones")]
        self.areas = [(str(p).strip().upper(), str(t).strip()) for p, t, *_ in filas("Áreas")]
        self.alias = {norm(a): str(b).strip() for a, b, *_ in filas("Alias") if a and b}
        self.atc_preferente = {clave(a): str(b).strip().upper() for a, b, *_ in filas("ATC preferente") if a and b}
        self.traducciones = {clave(a): str(b).strip() for a, b, *_ in filas("Traducciones") if a and b}
        self.correcciones = {}
        for nombre, columna, valor, *_ in filas("Correcciones"):
            if nombre and columna:
                self.correcciones.setdefault(norm(nombre), {})[str(columna).strip()] = "" if valor is None else str(valor)

    def es_cv(self, atc: str) -> bool:
        return bool(atc) and any(atc.startswith(p) for p in self.cv_incluidos) and not any(
            atc.startswith(p) for p in self.cv_excluidos
        )

    def descripcion(self, atc: str) -> str:
        return _prefijo_mas_largo(atc, self.descripciones) if atc else ""

    def area(self, atc: str) -> str:
        return _prefijo_mas_largo(atc, self.areas) if atc else ""

    def no_descomponer(self, atc: str) -> bool:
        return any(atc.startswith(p) for p in self.sin_descomponer) if atc else False

    def clasificar(self, ensayo: Ensayo, atcs: list[str], vias: list[str]) -> str:
        """Categorías del ensayo para una lista de ATC (uno por principio activo)."""
        elegidas = set()
        otros = False
        for atc in atcs:
            coinciden = [c.nombre for c in ensayo.categorias if c.coincide(atc, vias)]
            elegidas.update(coinciden)
            if not coinciden and self.es_cv(atc):
                otros = True
        valores = [c.nombre for c in ensayo.categorias if c.nombre in elegidas]
        valores = list(dict.fromkeys(valores))
        if otros:
            valores.append(ensayo.otros)
        return ", ".join(valores) if valores else ensayo.ninguna
